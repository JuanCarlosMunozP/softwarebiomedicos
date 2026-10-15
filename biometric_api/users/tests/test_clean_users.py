import csv
from io import StringIO

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError

from users.models import User


def _write(path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def _run(src, dest, *args):
    out = StringIO()
    err = StringIO()
    call_command(
        "clean_users",
        str(src),
        "-o",
        str(dest),
        *args,
        stdout=out,
        stderr=err,
    )
    return out.getvalue() + err.getvalue()


def _rows(path):
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def test_limpia_columnas_sin_imprimir_celdas(tmp_path):
    src = tmp_path / "usuarios.csv"
    _write(
        src,
        "username;email;password;first_name;last_name;phone;area;estado\n"
        "  user.uno  ; Uno@Example.com ;  clave secreta  ; Ana   Maria ;NO REGISTRA;"
        "300-123-4567abc;  UCI  ;inactivo\n"
        "user.dos;dos@example.com;clave2;Luis;Perez;no-es-telefono;Radiologia;\n"
        "user.uno;otro@example.com;clave3;Otra;Persona;3001234567;UCI;Activo\n"
        ";sin-user@example.com;clave4;A;B;3001234567;UCI;Activo\n"
        "user.tres;no-es-correo;clave5;A;B;3001234567;UCI;Activo\n"
        "user.cuatro;uno@example.com;clave6;A;B;3001234567;UCI;Activo\n",
    )
    dest = tmp_path / "salida.csv"
    report = _run(src, dest)

    assert "clave secreta" not in report
    assert "uno@example.com" not in report
    assert "user.uno" not in report
    assert "300-123-4567" not in report
    assert "otro@example.com" not in report
    assert "tecnico" in report
    assert "Listas para importar: 4" in report
    assert "Usernames con sufijo (el mismo username tenía otro correo): 1" in report
    assert "Usernames tomados del correo (venían vacíos): 1" in report

    rows = _rows(dest)
    assert [row["username"] for row in rows] == [
        "user.uno",
        "user.dos",
        "user.uno_2",
        "sin-user",
    ]
    assert rows[2]["email"] == "otro@example.com"
    assert rows[2]["password"] == "clave3"
    assert rows[3]["email"] == "sin-user@example.com"
    assert rows[0]["email"] == "uno@example.com"
    assert rows[0]["password"] == "clave secreta"
    assert rows[0]["role"] == "tecnico"
    assert rows[0]["first_name"] == "Ana Maria"
    assert rows[0]["last_name"] == ""
    assert rows[0]["phone"] == "300-123-4567"
    assert rows[0]["area"] == "UCI"
    assert rows[0]["estado"] == "Inactivo"
    assert rows[1]["phone"] == ""
    assert rows[1]["estado"] == "Activo"
    assert list(rows[0]) == [
        "username",
        "email",
        "password",
        "role",
        "first_name",
        "last_name",
        "phone",
        "area",
        "estado",
    ]


def test_rol_por_fila_y_valor_por_defecto(tmp_path):
    src = tmp_path / "usuarios.csv"
    _write(
        src,
        "username,email,password,role,estado\n"
        "u1,a@example.com,secret,operativo,Activo\n"
        "u2,b@example.com,secret,ingeniero_biomedico,1\n"
        "u3,c@example.com,secret,,Activo\n"
        "u4,d@example.com,secret,desconocido,Activo\n",
    )
    dest = tmp_path / "salida.csv"
    _run(src, dest, "--role", "admin")

    rows = _rows(dest)
    assert [(row["username"], row["role"]) for row in rows] == [
        ("u1", "tecnico"),
        ("u2", "ingeniero"),
        ("u3", "admin"),
    ]
    assert rows[1]["estado"] == "Activo"


def test_convierte_xlsx_a_csv_sin_imprimir_celdas(tmp_path):
    from openpyxl import Workbook

    src = tmp_path / "usuarios.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Usuarios"
    sheet.append(
        ["username", "email", "password", "first_name", "last_name", "phone", "area", "estado"]
    )
    sheet.append(
        [
            "user.uno",
            "Uno@Example.com",
            "clave secreta",
            "Ana",
            "Perez",
            3001234567.0,
            "UCI",
            "Inactivo",
        ]
    )
    book.save(src)
    book.close()

    dest = tmp_path / "salida.csv"
    report = _run(src, dest)

    assert "clave secreta" not in report
    assert "user.uno" not in report
    assert "3001234567" not in report
    assert "Uno@Example.com" not in report

    rows = _rows(dest)
    assert len(rows) == 1
    assert rows[0]["username"] == "user.uno"
    assert rows[0]["email"] == "uno@example.com"
    assert rows[0]["password"] == "clave secreta"
    assert rows[0]["phone"] == "3001234567"
    assert rows[0]["first_name"] == "Ana"
    assert rows[0]["area"] == "UCI"
    assert rows[0]["estado"] == "Inactivo"
    assert rows[0]["role"] == "tecnico"


def test_xlsx_repite_username_y_lee_mailto(tmp_path):
    from openpyxl import Workbook

    src = tmp_path / "usuarios.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(
        ["username", "email", "password", "first_name", "last_name", "phone", "area", "estado"]
    )
    sheet.append(["misma", "una@example.com", "clave-a", "Ana", "Uno", None, "UCI", "Activo"])
    sheet.append(
        ["misma", "Nombre <dos@example.com>", "clave-b", "Luis", "Dos", None, "UCI", "Activo"]
    )
    sheet.append(["NO REGISTRA", "ver ficha", "clave-c", "Mara", "Tres", None, "UCI", "Activo"])
    sheet.cell(4, 2).hyperlink = "mailto:mara@example.com"
    sheet.cell(4, 2).value = "ver ficha"
    book.save(src)
    book.close()

    dest = tmp_path / "salida.csv"
    report = _run(src, dest)

    assert "clave-a" not in report
    assert "mara@example.com" not in report
    assert "dos@example.com" not in report
    rows = _rows(dest)
    assert [(row["username"], row["email"]) for row in rows] == [
        ("misma", "una@example.com"),
        ("misma_2", "dos@example.com"),
        ("mara", "mara@example.com"),
    ]
    assert rows[2]["password"] == "clave-c"


def test_dry_run_no_escribe_archivo(tmp_path):
    src = tmp_path / "usuarios.csv"
    _write(src, "username,email,password\nuser.uno,a@example.com,secret\n")
    dest = tmp_path / "salida.csv"
    out = StringIO()
    call_command("clean_users", str(src), "--dry-run", stdout=out)
    assert not dest.exists()
    assert "DRY-RUN" in out.getvalue()
    assert "secret" not in out.getvalue()


def test_rechaza_rol_salida_igual_al_origen_y_cabecera_faltante(tmp_path):
    src = tmp_path / "usuarios.csv"
    _write(src, "username,email,password\nuser.uno,a@example.com,secret\n")
    with pytest.raises(CommandError):
        call_command("clean_users", str(src), "-o", str(src))

    bad = tmp_path / "sin_email.csv"
    _write(bad, "username,password\nuser.uno,secret\n")
    with pytest.raises(CommandError, match="email"):
        call_command("clean_users", str(bad), "-o", str(tmp_path / "salida.csv"))

    with pytest.raises(CommandError, match="--role"):
        call_command(
            "clean_users",
            str(src),
            "-o",
            str(tmp_path / "otra.csv"),
            "--role",
            "jefe",
        )


@pytest.mark.django_db
def test_el_archivo_limpio_entra_en_import_users(tmp_path):
    src = tmp_path / "usuarios.csv"
    _write(
        src,
        "username,email,password,first_name,last_name,phone,area,estado\n"
        "user.uno,Uno@Example.com,clave secreta,Ana,Perez,3001234567,UCI,Inactivo\n",
    )
    dest = tmp_path / "salida.csv"
    _run(src, dest)
    call_command("import_users", str(dest))

    user = User.objects.get(username="user.uno")
    assert user.email == "uno@example.com"
    assert user.check_password("clave secreta")
    assert user.role == User.Role.TECNICO
    assert user.first_name == "Ana"
    assert user.area == "UCI"
    assert user.is_active is False
